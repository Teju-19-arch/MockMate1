import io
import os

from pypdf import PdfReader
from docx import Document


ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx"
}

MAX_RESUME_SIZE = 5 * 1024 * 1024


def extract_resume_text(file_storage):
    """
    Extract text from PDF or DOCX resume.
    """

    if not file_storage:
        raise ValueError(
            "Please upload your resume."
        )

    filename = (
        file_storage.filename or ""
    ).strip()

    if not filename:
        raise ValueError(
            "Resume filename is missing."
        )

    extension = os.path.splitext(
        filename
    )[1].lower()

    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError(
            "Only PDF and DOCX resumes are supported."
        )

    file_bytes = file_storage.read()

    if not file_bytes:
        raise ValueError(
            "The uploaded resume is empty."
        )

    if len(file_bytes) > MAX_RESUME_SIZE:
        raise ValueError(
            "Resume must be smaller than 5 MB."
        )

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------

    if extension == ".pdf":

        try:

            reader = PdfReader(
                io.BytesIO(file_bytes)
            )

            pages = []

            for page in reader.pages:

                page_text = page.extract_text()

                if page_text:
                    pages.append(page_text)

            text = "\n".join(pages)

        except Exception as e:

            raise ValueError(
                f"Unable to read PDF resume: {e}"
            )

    # ------------------------------------------------------------------
    # DOCX
    # ------------------------------------------------------------------

    else:

        try:

            document = Document(
                io.BytesIO(file_bytes)
            )

            parts = []

            # Normal paragraphs
            for paragraph in document.paragraphs:

                value = paragraph.text.strip()

                if value:
                    parts.append(value)

            # Tables
            for table in document.tables:

                for row in table.rows:

                    cells = []

                    for cell in row.cells:

                        value = cell.text.strip()

                        if value:
                            cells.append(value)

                    if cells:
                        parts.append(
                            " | ".join(cells)
                        )

            text = "\n".join(parts)

        except Exception as e:

            raise ValueError(
                f"Unable to read DOCX resume: {e}"
            )

    text = text.strip()

    if not text:
        raise ValueError(
            "No readable text was found in the resume."
        )

    return text